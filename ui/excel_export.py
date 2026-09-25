"""Export the base-case DCF as an Excel workbook built from live formulas.

Every step is a formula referencing the Inputs sheet, so a reader can change
growth, margins, ROIC or the discount rate in Excel and watch the value update.
It mirrors engine/dcf.py and engine/projection.py; tests check that both give
the same value.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from engine.valuation import Assumptions

FONT = "Arial"
BLUE, BLACK, GREEN = "0000FF", "000000", "008000"
KEY_FILL = PatternFill("solid", fgColor="FFFF00")
HEAD_FILL = PatternFill("solid", fgColor="D9E1F2")
MONEY = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
MULT = '0.0"x"'
MAX_FADE = 20

# (row, label, key, number format, is key assumption)
_INPUTS = [
    (5, "Base-year revenue", "base_revenue", MONEY, False),
    (6, "Revenue growth, year 1", "growth_y1", PCT, True),
    (7, "Revenue growth, year 2", "growth_y2", PCT, True),
    (8, "Revenue growth, year 3", "growth_y3", PCT, False),
    (9, "Revenue growth, year 4", "growth_y4", PCT, False),
    (10, "Revenue growth, year 5", "growth_y5", PCT, True),
    (11, "EBIT margin, latest year", "ebit_margin", PCT, False),
    (12, "Target EBIT margin, year 5", "target_ebit_margin", PCT, True),
    (13, "Tax rate", "tax_rate", PCT, False),
    (14, "ROIC (return on new capital)", "roic", PCT, True),
    (15, "Fade period (years)", "fade_years", "0", True),
    (16, "Lasting excess return on new capital", "terminal_excess_return", PCT, False),
    (17, "Terminal growth", "terminal_growth", PCT, False),
    (18, "Discount rate (required return)", "discount_rate", PCT, True),
    (19, "Years since base fiscal year end", "years_since_fy_end", "0.00", False),
    (20, "Mid-year convention (years)", "mid_year", "0.0", False),
    (21, "Net debt", "net_debt", MONEY, False),
    (22, "Minority interest", "minority_interest", MONEY, False),
    (23, "Diluted shares", "shares_diluted", "#,##0", False),
    (24, "Share price", "price", "#,##0.00", False),
    (25, "Margin of safety", "margin_of_safety", PCT, False),
]
ROW = {key: row for row, _, key, _, _ in _INPUTS}
I = {key: f"Inputs!$B${row}" for key, row in ROW.items()}


def _font(color=BLACK, bold=False) -> Font:
    return Font(name=FONT, color=color, bold=bold)


def build_dcf_workbook(company: dict, a: Assumptions, notes: dict[str, str]) -> bytes:
    """`company`: name, ticker, currency, base_year, base_revenue, net_debt,
    minority_interest, shares_diluted, price, years_since_fy_end, data_source.
    `notes`: where each assumption came from, shown as cell comments."""
    path = a.growth_path()
    values = {**company, **a.__dict__, "mid_year": 0.5,
              "growth_y1": path[0], "growth_y2": path[1], "growth_y5": path[4]}
    if a.growth_override is not None:  # a segment build sets every year explicitly
        values.update(growth_y3=path[2], growth_y4=path[3])
    else:  # otherwise years 3-4 are the straight line from year 2 to year 5
        for key, n in (("growth_y3", 1), ("growth_y4", 2)):
            values[key] = f"=B{ROW['growth_y2']}+(B{ROW['growth_y5']}-B{ROW['growth_y2']})*{n}/3"
    wb = Workbook()
    _inputs_sheet(wb.active, company, values, notes)
    _dcf_sheet(wb.create_sheet("DCF"), company)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.font.name != FONT:
                    cell.font = Font(name=FONT, color=cell.font.color, bold=cell.font.bold)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _inputs_sheet(ws, company: dict, values: dict, notes: dict[str, str]) -> None:
    ws.title = "Inputs"
    ccy = company["currency"]
    ws["A1"] = f"{company['name']} ({company['ticker']}): base-case DCF inputs"
    ws["A1"].font = _font(bold=True)
    ws["A2"] = (f"Blue = input you can change; yellow = key assumptions. Money in {ccy}, "
                f"base year FY{company['base_year']}. Data: {company['data_source']}.")
    ws["A4"], ws["B4"], ws["C4"] = "Input", "Value", "Source / note"
    for c in ("A4", "B4", "C4"):
        ws[c].font, ws[c].fill = _font(bold=True), HEAD_FILL
    for row, label, key, fmt, is_key in _INPUTS:
        ws.cell(row, 1, label).font = _font()
        cell = ws.cell(row, 2, values[key])
        is_formula = isinstance(values[key], str) and values[key].startswith("=")
        cell.font, cell.number_format = _font(BLACK if is_formula else BLUE), fmt
        if is_key:
            cell.fill = KEY_FILL
        note = notes.get(key, "")
        ws.cell(row, 3, note).font = _font()
        if note:
            cell.comment = Comment(note, "valuation-dashboard")
    ws.cell(ROW["mid_year"], 3, "Cash flows arrive through the year, so each is discounted half a year less")
    ws.cell(ROW["years_since_fy_end"], 3, "The valuation is as of today, not the last fiscal year end")
    if not notes.get("growth_y3"):
        for key in ("growth_y3", "growth_y4"):
            ws.cell(ROW[key], 3, "Formula: straight line from year 2 to year 5")
    fade = DataValidation(type="list", formula1='"5,10,15,20"', allow_blank=False)
    ws.add_data_validation(fade)
    fade.add(f"B{ROW['fade_years']}")
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 90


def _dcf_sheet(ws, company: dict) -> None:
    ccy = company["currency"]
    ws["A1"] = f"{company['name']} ({company['ticker']}): base-case DCF ({ccy})"
    ws["A1"].font = _font(bold=True)
    ws["A2"] = ("Reinvestment = NOPAT x growth / ROIC. Stage 1: explicit years. Stage 2: growth fades "
                "to terminal growth and ROIC to the discount rate plus any lasting excess return. "
                "Fade rows beyond the fade period are switched off.")
    heads = ["Year", "Stage", "Revenue growth", f"Revenue ({ccy})", "EBIT margin", f"NOPAT ({ccy})",
             "ROIC", "Reinvestment rate", f"Reinvestment ({ccy})", f"Free cash flow ({ccy})",
             "Discount period", "Discount factor", f"PV of FCF ({ccy})", "In use"]
    for col, h in enumerate(heads, start=1):
        c = ws.cell(4, col, h)
        c.font, c.fill = _font(bold=True), HEAD_FILL
        c.alignment = Alignment(wrap_text=True, horizontal="center")

    g, t, r = I["tax_rate"], I["terminal_growth"], I["discount_rate"]
    shift = f"({I['mid_year']}+{I['years_since_fy_end']})"
    ws["A5"], ws["B5"] = f"FY{company['base_year']}", "Base"
    ws["D5"] = f"={I['base_revenue']}"
    ws["E5"] = f"={I['ebit_margin']}"
    ws["F5"] = f"=D5*E5*(1-{g})"
    ws["D5"].font = ws["E5"].font = _font(GREEN)

    for i in range(1, 6):  # Stage 1
        row = 5 + i
        ws.cell(row, 1, str(i))
        ws.cell(row, 2, "Stage 1")
        ws.cell(row, 3, f"={I[f'growth_y{i}']}").font = _font(GREEN)
        ws.cell(row, 4, f"=D{row - 1}*(1+C{row})")
        ws.cell(row, 5, f"={I['ebit_margin']}+({I['target_ebit_margin']}-{I['ebit_margin']})*{i}/5")
        ws.cell(row, 6, f"=D{row}*E{row}*(1-{g})")
        ws.cell(row, 7, f"={I['roic']}").font = _font(GREEN)
        _cash_flow_cells(ws, row, i, shift, r, in_use=1)

    for i in range(1, MAX_FADE + 1):  # Stage 2: the fade
        row = 10 + i
        frac = f"MIN({i},{I['fade_years']})/{I['fade_years']}"
        ws.cell(row, 1, str(5 + i))
        ws.cell(row, 2, "Fade")
        ws.cell(row, 3, f"={I['growth_y5']}+({t}-{I['growth_y5']})*{frac}")
        ws.cell(row, 4, f"=D{row - 1}*(1+C{row})")
        ws.cell(row, 5, f"=E{row - 1}")
        ws.cell(row, 6, f"=F{row - 1}*(1+C{row})")
        ws.cell(row, 7, f"={I['roic']}+(({r}+{I['terminal_excess_return']})-{I['roic']})*{frac}")
        _cash_flow_cells(ws, row, 5 + i, shift, r, in_use=f"=IF({i}<={I['fade_years']},1,0)")

    last = 10 + MAX_FADE
    rows = [
        ("Terminal value", None, None),
        ("Final-year NOPAT (end of fade)", f"=INDEX(F11:F{last},{I['fade_years']})", MONEY),
        ("Terminal return on new capital", f"={r}+{I['terminal_excess_return']}", PCT),
        ("Terminal free cash flow", "=B{nopat}*(1+{t})*(1-{t}/B{ronic})", MONEY),
        ("Terminal value", "=B{tfcf}/({r}-{t})", MONEY),
        ("Discount period", f"=5+{I['fade_years']}-{shift}", "0.00"),
        ("PV of terminal value", "=B{tv}/(1+{r})^B{period}", MONEY),
        ("Valuation", None, None),
        ("PV of Stage 1", "=SUM(M6:M10)", MONEY),
        ("PV of fade", f"=SUM(M11:M{last})", MONEY),
        ("PV of terminal value", "=B{pvtv}", MONEY),
        ("Enterprise value", "=B{pv1}+B{pv2}+B{pv3}", MONEY),
        ("Less net debt", f"=-{I['net_debt']}", MONEY),
        ("Less minority interest", f"=-{I['minority_interest']}", MONEY),
        ("Equity value", "=B{ev}+B{nd}+B{mi}", MONEY),
        ("Value per share (base case)", f"=B{{eq}}/{I['shares_diluted']}", "#,##0.00"),
        ("Share price", f"={I['price']}", "#,##0.00"),
        ("Upside / (downside)", "=B{vps}/B{px}-1", PCT),
        ("Buy below (after margin of safety)", f"=B{{vps}}*(1-{I['margin_of_safety']})", "#,##0.00"),
        ("Terminal value share of EV", "=B{pv3}/B{ev}", PCT),
        ("Implied exit multiple (TV / final NOPAT)", "=B{tv}/B{nopat}", MULT),
    ]
    start = last + 2
    at = {}
    keys = [None, "nopat", "ronic", "tfcf", "tv", "period", "pvtv", None, "pv1", "pv2", "pv3", "ev",
            "nd", "mi", "eq", "vps", "px", None, None, None, None]
    for n, key in enumerate(keys):
        if key:
            at[key] = start + n
    for n, (label, formula, fmt) in enumerate(rows):
        row = start + n
        cell = ws.cell(row, 1, label)
        if formula is None:
            cell.font = _font(bold=True)
            cell.fill = HEAD_FILL
            continue
        cell.font = _font(bold=label in ("Enterprise value", "Equity value", "Value per share (base case)"))
        v = ws.cell(row, 2, formula.format(t=t, r=r, **at))
        v.number_format = fmt
        if label == "Value per share (base case)":
            v.font, v.border = _font(bold=True), Border(top=Side("thin"), bottom=Side("double"))

    widths = [34, 10, 12, 16, 11, 16, 10, 12, 16, 16, 11, 11, 16, 8]
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + col)].width = w
    ws.freeze_panes = "C5"


def _cash_flow_cells(ws, row: int, year: int, shift: str, r: str, in_use: int | str) -> None:
    ws.cell(row, 8, f"=C{row}/G{row}")
    ws.cell(row, 9, f"=F{row}*H{row}")
    ws.cell(row, 10, f"=F{row}-I{row}")
    ws.cell(row, 11, f"={year}-{shift}")
    ws.cell(row, 12, f"=1/(1+{r})^K{row}")
    ws.cell(row, 13, f"=J{row}*L{row}*N{row}")
    ws.cell(row, 14, in_use)
    for col, fmt in [(3, PCT), (4, MONEY), (5, PCT), (6, MONEY), (7, PCT), (8, PCT), (9, MONEY),
                     (10, MONEY), (11, "0.00"), (12, "0.000"), (13, MONEY), (14, "0")]:
        ws.cell(row, col).number_format = fmt
